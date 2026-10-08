from __future__ import annotations

import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import database
from auth import hash_password, is_valid_email, validate_password, verify_password
from ingestion import load_documents_from_bytes
from pptx import Presentation
from pptx.util import Inches
from io import BytesIO


class PasswordTests(unittest.TestCase):
    def test_password_hash_is_salted_and_verifiable(self) -> None:
        salt, password_hash = hash_password("a secure example password")
        other_salt, other_hash = hash_password("a secure example password")
        self.assertNotEqual(salt, other_salt)
        self.assertNotEqual(password_hash, other_hash)
        self.assertTrue(verify_password("a secure example password", salt, password_hash))
        self.assertFalse(verify_password("a different password", salt, password_hash))

    def test_email_and_password_validation(self) -> None:
        self.assertTrue(is_valid_email("Reader@example.org"))
        self.assertFalse(is_valid_email("not-an-email"))
        self.assertIsNotNone(validate_password("short"))
        self.assertIsNone(validate_password("long-enough-example-password"))


class DatabaseIsolationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.env = patch.dict(
            "os.environ",
            {"NOTES_NAVIGATOR_DB_PATH": str(Path(self.temp_dir.name) / "test.sqlite3")},
        )
        self.env.start()
        database.initialize_database()
        salt_a, hash_a = hash_password("user a secure password")
        salt_b, hash_b = hash_password("user b secure password")
        self.user_a = database.create_account("a@example.org", salt_a, hash_a)
        self.user_b = database.create_account("b@example.org", salt_b, hash_b)

    def tearDown(self) -> None:
        self.env.stop()
        self.temp_dir.cleanup()

    def test_accounts_get_separate_spaces_and_chats(self) -> None:
        self.assertNotEqual(self.user_a["space_id"], self.user_b["space_id"])
        self.assertEqual(
            database.list_chats(self.user_a["id"], self.user_a["space_id"])[0]["id"],
            self.user_a["chat_id"],
        )
        self.assertEqual(
            database.list_chats(self.user_b["id"], self.user_a["space_id"]),
            [],
        )

    def test_documents_are_scoped_by_user_and_space(self) -> None:
        object_path = (
            f"/objects/{self.user_a['id']}/{self.user_a['space_id']}/uploads/test-object"
        )
        database.add_document(
            document_id="document-a",
            user_id=self.user_a["id"],
            space_id=self.user_a["space_id"],
            file_name="notes.txt",
            object_path=object_path,
            sha256="f" * 64,
            byte_size=8,
        )
        self.assertEqual(
            database.list_documents(self.user_b["id"], self.user_a["space_id"]),
            [],
        )
        self.assertEqual(
            len(database.list_documents(self.user_a["id"], self.user_a["space_id"])),
            1,
        )

    def test_foreign_user_cannot_create_chat_in_another_users_space(self) -> None:
        with self.assertRaises(sqlite3.IntegrityError):
            database.create_chat(self.user_b["id"], self.user_a["space_id"])
        self.assertFalse(
            database.delete_chat(
                self.user_b["id"],
                self.user_a["space_id"],
                self.user_a["chat_id"],
            )
        )

    def test_deleting_space_cascades_chat_and_document_metadata(self) -> None:
        object_path = (
            f"/objects/{self.user_a['id']}/{self.user_a['space_id']}/uploads/test-object"
        )
        database.add_document(
            document_id="document-a",
            user_id=self.user_a["id"],
            space_id=self.user_a["space_id"],
            file_name="notes.txt",
            object_path=object_path,
            sha256="e" * 64,
            byte_size=8,
        )
        self.assertTrue(database.delete_space(self.user_a["id"], self.user_a["space_id"]))
        self.assertEqual(database.list_documents(self.user_a["id"], self.user_a["space_id"]), [])
        self.assertEqual(database.list_chats(self.user_a["id"], self.user_a["space_id"]), [])


class PowerPointIngestionTests(unittest.TestCase):
    def test_pptx_extracts_slide_text_and_source_metadata(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        textbox.text = "Photosynthesis converts light into chemical energy."
        buffer = BytesIO()
        presentation.save(buffer)

        documents = load_documents_from_bytes("lesson.pptx", buffer.getvalue())
        self.assertEqual(len(documents), 1)
        self.assertIn("Photosynthesis", documents[0].page_content)
        self.assertEqual(documents[0].metadata["slide"], 1)
        self.assertEqual(documents[0].metadata["file_name"], "lesson.pptx")

    def test_text_ingestion_adds_file_metadata(self) -> None:
        documents = load_documents_from_bytes("lesson.txt", b"Earth orbits the Sun.")
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].metadata["file_name"], "lesson.txt")
        self.assertIn("Earth", documents[0].page_content)


if __name__ == "__main__":
    unittest.main()
