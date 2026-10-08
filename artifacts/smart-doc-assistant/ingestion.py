from __future__ import annotations

import shutil
import subprocess
import tempfile
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation


EMBEDDING_MODEL = "models/gemini-embedding-001"
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150


def _add_source_metadata(
    documents: list[Document],
    file_name: str,
) -> list[Document]:
    for document in documents:
        document.metadata["source"] = file_name
        document.metadata["file_name"] = file_name
    return documents


def _load_legacy_powerpoint(content: bytes, file_name: str) -> list[Document]:
    converter = shutil.which("libreoffice") or shutil.which("soffice")
    if not converter:
        raise RuntimeError(
            "Legacy .ppt files require LibreOffice conversion, which is not available."
        )

    with tempfile.TemporaryDirectory(prefix="notes-navigator-ppt-") as temp_dir:
        source_path = Path(temp_dir) / "legacy.ppt"
        source_path.write_bytes(content)
        profile_path = Path(temp_dir) / "lo-profile"
        try:
            subprocess.run(
                [
                    converter,
                    "--headless",
                    f"-env:UserInstallation={profile_path.as_uri()}",
                    "--convert-to",
                    "pptx",
                    "--outdir",
                    temp_dir,
                    str(source_path),
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=90,
            )
        except subprocess.TimeoutExpired as error:
            raise RuntimeError(f"Conversion timed out for {file_name}.") from error
        except subprocess.CalledProcessError as error:
            detail = (error.stderr or error.stdout or "").strip()
            raise RuntimeError(
                f"Could not convert {file_name} to a readable PowerPoint file."
                + (f" {detail[:300]}" if detail else "")
            ) from error

        converted_path = Path(temp_dir) / "legacy.pptx"
        if not converted_path.is_file():
            raise RuntimeError(f"Could not convert {file_name} to .pptx.")
        return _load_powerpoint(converted_path.read_bytes(), file_name)


def _load_powerpoint(content: bytes, file_name: str) -> list[Document]:
    try:
        presentation = Presentation(BytesIO(content))
    except Exception as error:
        raise RuntimeError(f"Could not read PowerPoint file {file_name}.") from error

    documents: list[Document] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    slide_parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cell_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if cell_text:
                        slide_parts.append(cell_text)

        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, KeyError):
            notes = ""
        if notes and notes not in slide_parts:
            slide_parts.append(f"Speaker notes: {notes}")

        content_text = "\n".join(slide_parts).strip()
        if content_text:
            documents.append(
                Document(
                    page_content=f"Slide {slide_number}\n{content_text}",
                    metadata={
                        "source": file_name,
                        "file_name": file_name,
                        "slide": slide_number,
                    },
                )
            )
    return documents


def load_documents_from_bytes(file_name: str, content: bytes) -> list[Document]:
    suffix = Path(file_name).suffix.lower()
    if suffix not in {".pdf", ".txt", ".ppt", ".pptx"}:
        raise ValueError(
            f"{file_name}: only PDF, TXT, PPT, and PPTX files are supported."
        )
    if not content:
        raise ValueError(f"{file_name} is empty.")

    if suffix == ".pptx":
        return _load_powerpoint(content, file_name)
    if suffix == ".ppt":
        return _load_legacy_powerpoint(content, file_name)

    with tempfile.TemporaryDirectory(prefix="notes-navigator-doc-") as temp_dir:
        temp_path = Path(temp_dir) / f"{uuid.uuid4().hex}{suffix}"
        temp_path.write_bytes(content)
        try:
            if suffix == ".pdf":
                documents = PyPDFLoader(str(temp_path)).load()
            else:
                documents = TextLoader(
                    str(temp_path),
                    encoding="utf-8",
                    autodetect_encoding=True,
                ).load()
        except Exception as error:
            raise RuntimeError(f"Could not read {file_name}: {error}") from error
    return _add_source_metadata(documents, file_name)


def split_documents(documents: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        add_start_index=True,
    )
    chunks = splitter.split_documents(documents)
    return [chunk for chunk in chunks if chunk.page_content.strip()]


def build_vector_store(
    documents: list[Document],
    *,
    api_key: str,
) -> tuple[FAISS, int]:
    chunks = split_documents(documents)
    if not chunks:
        raise ValueError("The selected files did not contain text to index.")
    embeddings = GoogleGenerativeAIEmbeddings(
        model=EMBEDDING_MODEL,
        api_key=api_key,
    )
    return FAISS.from_documents(chunks, embeddings), len(chunks)


def serialize_vector_store(vector_store: FAISS) -> bytes:
    return vector_store.serialize_to_bytes()


def restore_vector_store(serialized: bytes, *, api_key: str) -> FAISS:
    embeddings = GoogleGenerativeAIEmbeddings(
        model=EMBEDDING_MODEL,
        api_key=api_key,
    )
    # These bytes are generated only by this app and stored in its private bucket.
    return FAISS.deserialize_from_bytes(
        serialized,
        embeddings,
        allow_dangerous_deserialization=True,
    )


def source_label(metadata: dict[str, Any]) -> str:
    filename = str(metadata.get("file_name") or metadata.get("source") or "Document")
    if metadata.get("slide") is not None:
        return f"{filename} · slide {metadata['slide']}"
    page = metadata.get("page")
    if page is not None:
        try:
            return f"{filename} · page {int(page) + 1}"
        except (TypeError, ValueError):
            return f"{filename} · page {page}"
    return filename
